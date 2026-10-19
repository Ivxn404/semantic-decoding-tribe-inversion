"""Fill the VIT Project-I poster template (Poster_-_Capstone_-_Design.pptx) with the final content.
Canvas 33.11 x 46.81 in (A0 proportions); printed on A3, so 26 pt here is ~9 pt on paper."""
import copy, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from PIL import Image

import os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TPL = os.environ.get("POSTER_TEMPLATE", "Poster_-_Capstone_-_Design.pptx")  # VIT template (not redistributed)
OUT = os.path.join(ROOT, "poster", "Poster_Sarkar_Choudhuri_Project1.pptx")
FIG = os.path.join(ROOT, "figures")
LOGO = os.environ.get("VIT_LOGO", "vit_logo.jpeg")  # taken from the VIT report template
QR = os.path.join(ROOT, "poster", "repo_qr_plain.png")

NAVY = RGBColor(0x13, 0x24, 0x3B); TEAL = RGBColor(0x0B, 0x5E, 0x66); INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55); ORANGE = RGBColor(0xC2, 0x4E, 0x14); TINT = RGBColor(0xE8, 0xF3, 0xF1)
GOLD = RGBColor(0xFF, 0xF1, 0xC4); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HF, BF = "Cambria", "Calibri"
H, BODY, SMALL = 40, 25, 21

prs = Presentation(TPL)
sl = prs.slides[0]
box = {sh.text_frame.text.strip(): sh for sh in sl.shapes if sh.has_text_frame}

L, R, W = 0.5, 16.76, 15.85
GEOM = {
    "VIT logo, School name, Program name": (0.5, 0.35, 32.11, 3.0),
    "Project title": (0.5, 3.65, 32.11, 3.0),
    "Team members and Guide": (L, 7.0, W, 5.2),
    "Project Overview": (R, 7.0, W, 5.2),
    "Justification of Complex Engineering Problem": (L, 12.6, W, 9.4),
    "Proposed Block Diagram": (R, 12.6, W, 9.4),
    "Design Methodology and Novelty": (L, 22.4, W, 13.6),
    "Project Results": (R, 22.4, W, 13.6),
    "PO and PSO mapping with justification": (0.5, 36.4, 10.44, 7.0),
    "WP mapping with justification": (11.33, 36.4, 10.44, 7.0),
    "SDG mapping with justification": (22.17, 36.4, 10.44, 7.0),
    "References": (L, 43.8, W, 2.7),
    "Project outcome": (R, 43.8, W, 2.7),
}


def place(sh, g):
    sh.left, sh.top, sh.width, sh.height = Inches(g[0]), Inches(g[1]), Inches(g[2]), Inches(g[3])


def runs(p, parts, size, font=BF, color=INK):
    """parts: str or list of (text, bold) tuples."""
    if isinstance(parts, str):
        parts = [(parts, False)]
    for t, b in parts:
        r = p.add_run(); r.text = t
        f = r.font; f.size = Pt(size); f.bold = b; f.name = font; f.color.rgb = color


def fill(sh, heading, paras, size=BODY, anchor=MSO_ANCHOR.TOP, head_size=H, bullets=False, space=6):
    tf = sh.text_frame
    tf.word_wrap = True
    bp = tf._txBody.bodyPr
    for child in list(bp):
        if child.tag.endswith("Autofit"):
            bp.remove(child)
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.25); tf.margin_top = Inches(0.15); tf.margin_bottom = Inches(0.1)
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    if heading:
        runs(p0, [(heading, True)], head_size, HF, TEAL)
        p0.space_after = Pt(8)
        first = False
    else:
        first = True
    for para in paras:
        p = p0 if first else tf.add_paragraph()
        first = False
        runs(p, para, size)
        p.space_after = Pt(space)
        p.line_spacing = 1.05
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.3))); pPr.set("indent", str(-Inches(0.25)))
            from pptx.oxml.ns import qn
            bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
            pPr.append(bu)
    return tf


def picture(path, x, y, w, h):
    """Fit image inside (x, y, w, h) inches, centred."""
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih); pw, ph = iw * s, ih * s
    return sl.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def rect(x, y, w, h, fill_rgb, line_rgb=None, text=None, size=22, bold=False, color=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill_rgb
    if line_rgb is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_rgb; s.line.width = Pt(2.5)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    if text is not None:
        tf = s.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.12)
        lines = text if isinstance(text, list) else [text]
        for k, ln in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            runs(p, [(ln, bold if k == 0 else False)], size if k == 0 else size - 4, BF, color)
    return s


def arrow(x1, y1, x2, y2, rgb=NAVY, wpt=4):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb; c.line.width = Pt(wpt)
    ln = c.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def table(x, y, w, rows, colw, size=19, header_rgb=NAVY, rowh=0.42, bold_last=False):
    t = sl.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows))).table
    for j, cw in enumerate(colw):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            c = t.cell(i, j); c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            hdr = i == 0
            last = bold_last and i == len(rows) - 1
            runs(p, [(val, hdr or last)], size, BF, WHITE if hdr else INK)
            c.fill.solid()
            c.fill.fore_color.rgb = header_rgb if hdr else (GOLD if last else (WHITE if i % 2 else TINT))
    return t


for k, g in GEOM.items():
    place(box[k], g)

# ---------------------------------------------------------------- header + title
hb = box["VIT logo, School name, Program name"]
fill(hb, None, [[("School of Electronics Engineering (SENSE)", True)],
                [("B.Tech. Electronics and Communication Engineering (specialisation in Biomedical Engineering)", False)],
                [("BECE497J Project-I  ·  Final Review and Poster Presentation  ·  29 October 2026", False)]],
     size=30, anchor=MSO_ANCHOR.MIDDLE, space=2)
for p in hb.text_frame.paragraphs:
    p.alignment = PP_ALIGN.CENTER
hb.text_frame.margin_left = Inches(7.0); hb.text_frame.margin_right = Inches(4.0)
picture(LOGO, 0.9, 0.6, 6.2, 2.5)
picture(QR, 29.0, 0.55, 2.6, 2.6)
qt = sl.shapes.add_textbox(Inches(26.0), Inches(1.2), Inches(2.95), Inches(1.4))
qt.text_frame.word_wrap = True
p = qt.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
runs(p, [("Code, data and demo", True)], 20, BF, TEAL)
p2 = qt.text_frame.add_paragraph(); p2.alignment = PP_ALIGN.RIGHT
runs(p2, "github.com/Ivxn404/semantic-decoding-tribe-inversion", 15, BF, GREY)

tb = box["Project title"]
fill(tb, None, [[("Conditioning and Non-Convexity of Semantic Decoding by Inversion of Deep Brain Encoding Models", True)],
                [("Predictable is not recoverable: inverting Meta's TRIBE v2 foundation brain encoder on synthetic and real fMRI", False)]],
     size=58, anchor=MSO_ANCHOR.MIDDLE, space=6)
tps = tb.text_frame.paragraphs
for p in tps:
    p.alignment = PP_ALIGN.CENTER
for r in tps[0].runs:
    r.font.name = HF; r.font.color.rgb = NAVY; r.font.size = Pt(62)
for r in tps[1].runs:
    r.font.size = Pt(32); r.font.color.rgb = TEAL; r.font.italic = True

# ---------------------------------------------------------------- team / overview
fill(box["Team members and Guide"], "Team members and Guide", [
    [("Abhranil Sarkar", True), ("  (23BML0077)  ·  computational track: model instrumentation, inversion engine, experiments, real-fMRI pipeline", False)],
    [("Tanisha Choudhuri", True), ("  (23BML0081)  ·  neurobiological track: cortical networks and ROI definition, dataset curation, interpretation", False)],
    [("Guide: ", True), ("Dr. Sharmila Nageswaran, Associate Professor, School of Electronics Engineering (SENSE), VIT Vellore", False)],
], size=29, bullets=True, space=10)

fill(box["Project Overview"], "Project Overview", [
    "Brain encoding models predict fMRI from what a person sees and hears; decoding runs the other way. We ask whether a near-perfect encoder can simply be inverted to read meaning back out of the brain.",
    [("We invert Meta's open 177 M-parameter TRIBE v2 encoder by gradient descent, measure ", False), ("uniqueness, conditioning and regional recoverability", True),
     (" with exact ground truth, and validate on real fMRI from a held-out story.", False)],
], size=29, space=10)

# ---------------------------------------------------------------- complex engineering problem
fill(box["Justification of Complex Engineering Problem"], "Justification of Complex Engineering Problem", [
    [("No closed-form answer. ", True), ("The forward model is a 177 M-parameter nonlinear transformer mapping a 32 × 1 152 representation to 20 484 cortical vertices; its inverse cannot be written down and must be measured.", False)],
    [("Ill-posed by nature. ", True), ("Existence, uniqueness and stability of the inverse (Hadamard) are all open: fMRI is slow, blurred and noisy.", False)],
    [("No ground truth in real data. ", True), ("The true meaning behind a brain recording is unknown, so we built a synthetic framework with exact error, then a real-fMRI test with identification at a known chance level.", False)],
    [("Engineering obstacle. ", True), ("The released pipeline is not differentiable (speech synthesis, transcription, frozen extractors); we found and re-entered the one differentiable boundary.", False)],
    [("Conflicting requirements. ", True), ("Decoding fidelity vs non-invasive, low-cost acquisition vs mental privacy cannot all be maximised.", False)],
    [("Interdisciplinary. ", True), ("Inverse-problem theory, deep learning, neuroimaging (surface registration, atlases) and neuroanatomy must work together.", False)],
], size=27, bullets=True, space=9)

# ---------------------------------------------------------------- block diagram (native shapes)
bx, by = R, 12.6
fill(box["Proposed Block Diagram"], "Proposed Block Diagram", [], size=BODY)
Y1, Y2, Y3 = by + 1.35, by + 4.05, by + 6.75
bw, bh, gx = 3.45, 1.85, 0.55
xs = [bx + 0.35 + k * (bw + gx) for k in range(4)]
rect(xs[0], Y1, bw, bh, TINT, TEAL, ["Stimulus", "story audio + words"], 22, True)
rect(xs[1], Y1, bw, bh, WHITE, GREY, ["Feature extractors", "w2v-BERT, Llama-3.2-3B"], 22, True)
rect(xs[2], Y1, bw, bh, GOLD, ORANGE, ["Representation x", "32 × 1 152 (free)"], 22, True)
rect(xs[3], Y1, bw, bh, NAVY, None, ["TRIBE v2, frozen", "177 M parameters"], 22, True, WHITE)
for k in range(3):
    arrow(xs[k] + bw, Y1 + bh / 2, xs[k + 1], Y1 + bh / 2)
rect(xs[3], Y2, bw, bh, TINT, TEAL, ["Predicted fMRI f(x)", "20 484 vertices"], 22, True)
arrow(xs[3] + bw / 2, Y1 + bh, xs[3] + bw / 2, Y2)
rect(xs[2], Y2, bw, bh, WHITE, GREY, ["Loss", "‖f(x) − y‖² or 1 − r"], 22, True)
arrow(xs[3], Y2 + bh / 2, xs[2] + bw, Y2 + bh / 2)
rect(xs[1], Y2, bw, bh, TINT, TEAL, ["Target fMRI y", "synthetic or real BOLD"], 22, True)
arrow(xs[1] + bw, Y2 + bh / 2, xs[2], Y2 + bh / 2)
rect(xs[0], Y2, bw, bh, GOLD, ORANGE, ["Adam update", "gradient to x only"], 22, True)
arrow(xs[2] + 0.4, Y2 + bh, xs[2] + 0.4, Y2 + bh + 0.35, ORANGE)
c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(xs[0] + bw / 2), Inches(Y2 + bh + 0.35), Inches(xs[2] + 0.4), Inches(Y2 + bh + 0.35)); c.line.color.rgb = ORANGE; c.line.width = Pt(4)
arrow(xs[0] + bw / 2, Y2 + bh + 0.35, xs[0] + bw / 2, Y2 + bh, ORANGE)
arrow(xs[0] + bw / 2, Y2, xs[2] + 0.6, Y1 + bh, ORANGE)
rect(xs[0], Y3 + 0.2, bw * 2 + gx, 1.5, NAVY, None, ["Evaluate x̂: cosine to truth, identification rank"], 22, True, WHITE)
rect(xs[2], Y3 + 0.2, bw * 2 + gx, 1.5, TEAL, None, ["Controls: inits × 4 truths, barrier test, ROI vs random"], 22, True, WHITE)
cap = sl.shapes.add_textbox(Inches(bx + 0.3), Inches(by + 8.55), Inches(W - 0.6), Inches(0.7))
cap.text_frame.word_wrap = True
runs(cap.text_frame.paragraphs[0], "Only x is optimised; network weights stay frozen. Re-entering after feature aggregation reproduces TRIBE v2's predictions exactly (r = 1.0000).", 19, BF, GREY)

# ---------------------------------------------------------------- methodology + novelty
mx, my = L, 22.4
fill(box["Design Methodology and Novelty"], "Design Methodology and Novelty", [
    [("1  Forward operator. ", True), ("TRIBE v2 (trained on 451.6 h of fMRI from 25 people watching films and video clips and listening to stories), frozen.", False)],
    [("2  Inversion. ", True), ("x̂ = argmin ‖f(x) − y‖², Adam (lr 0.05, cosine schedule), weights frozen.", False)],
    [("3  Exact evaluation. ", True), ("Known synthetic x → y = f(x) → recover x; budget, noise, extent, priors, 8 initialisations × 4 truths.", False)],
    [("4  Mechanism tests. ", True), ("Forward-only barrier test along solution differences; Jacobian sensitivity vs 200 random directions.", False)],
    [("5  Regions and real data. ", True), ("6 Destrieux networks vs size-matched random sets; LeBel 2023 fMRI projected to fsaverage5, held-out story.", False)],
    [("Why fMRI, not EEG: ", True), ("millimetre detail on 20 484 vertices, the only deep encoder large enough to invert, and one inverse problem instead of two (EEG adds source localisation). EEG is faster and cheaper, which matters for a device, not for this measurement.", False)],
], size=26, space=8)
table(mx + 0.3, my + 8.25, W - 0.6, [
    ["Novelty: closest prior work", "Deep encoder", "Optimises input", "Tests uniqueness", "Per region"],
    ["Kay et al. 2008", "○", "○", "○", "○"],
    ["Gardner & Liu 2019 (linear IEM)", "○", "●", "●", "○"],
    ["Tang et al. 2023", "○", "●", "○", "●"],
    ["TRIBE v2 (2026) and its two follow-ups", "●", "○ / ●", "○", "○"],
    ["This work", "●", "●", "●", "●"],
], [6.05, 2.2, 2.3, 2.5, 2.2], size=19, rowh=0.6, bold_last=True)
nt = sl.shapes.add_textbox(Inches(mx + 0.3), Inches(my + 12.05), Inches(W - 0.6), Inches(1.4))
nt.text_frame.word_wrap = True
runs(nt.text_frame.paragraphs[0], "Literature search (9 Oct 2026, arXiv / Google Scholar / web): no published uniqueness or conditioning analysis of a deep foundation brain encoder.", 19, BF, GREY)

# ---------------------------------------------------------------- results
rx, ry = R, 22.4
fill(box["Project Results"], "Project Results", [
    [("A good fit is not a good decode, and the inverse identifies but does not reconstruct.", True)]], size=23)
picture(f"{FIG}/fig_dissociation.png", rx + 0.25, ry + 1.75, 7.6, 3.8)
picture(f"{FIG}/fig_nullspace_test.png", rx + 7.95, ry + 1.75, 7.6, 3.8)
picture(f"{FIG}/fig_roi_recovery.png", rx + 0.3, ry + 5.65, W - 0.6, 3.9)
picture(f"{FIG}/fig_realdata_audio_text.png", rx + 0.3, ry + 9.6, W - 0.6, 3.15)
kt = sl.shapes.add_textbox(Inches(rx + 0.3), Inches(ry + 12.75), Inches(W - 0.6), Inches(0.8))
kt.text_frame.word_wrap = True
runs(kt.text_frame.paragraphs[0], [("Synthetic: ", True), ("error ↓28.4× but recovery saturates (0.944); equal fits disagree (cosine 0.817); barrier, not null space. ", False),
                                   ("Regions: ", True), ("visual −0.125, somatomotor +0.102 vs random. ", False),
                                   ("Real fMRI: ", True), ("r = 0.48 (56% of ceiling); 4/5 segments identified (chance 2.2%), cosine 0.02.", False)], 17, BF, INK)

# ---------------------------------------------------------------- PO / WP / SDG
fill(box["PO and PSO mapping with justification"], "PO and PSO mapping", [
    [("PO1 ", True), ("Linear algebra, optimisation and inverse-problem theory applied to a neural operator.", False)],
    [("PO2 ", True), ("Formulated decoding as an ill-posed inverse problem and analysed its causes.", False)],
    [("PO4 ", True), ("Designed controlled experiments with exact ground truth and matched controls.", False)],
    [("PO5 ", True), ("PyTorch, TRIBE v2, ANTs, nilearn, GPU (Colab) computing.", False)],
    [("PO6, PO8 ", True), ("Mental-privacy and clinical-use implications; open data, licences respected.", False)],
    [("PO9, PO10 ", True), ("Two-track teamwork; IEEE manuscript, poster, reviews.", False)],
    [("PO12 ", True), ("Self-taught foundation models and neuroimaging tools.", False)],
    [("PSO2, PSO3 ", True), ("Analysis and imaging of biosignals (fMRI) with modern tools.", False)],
], size=24, space=6)
fill(box["WP mapping with justification"], "WP mapping", [
    [("WP1 Depth of knowledge: ", True), ("neuroscience, fMRI, deep learning and inverse-problem theory together.", False)],
    [("WP2 Conflicting requirements: ", True), ("fidelity vs non-invasiveness vs privacy.", False)],
    [("WP3 Depth of analysis: ", True), ("separates non-convexity from information loss with falsifiable tests.", False)],
    [("WP4 Familiarity of issues: ", True), ("slow, noisy, blurred BOLD; no ground truth in real data.", False)],
    [("WP7 Interdependence: ", True), ("ROI anatomy and computation depend on each other.", False)],
    [("Knowledge profile: ", True), ("WK3 signals and systems, WK4 inverse problems, WK5 design of the inversion pipeline, WK6 modern tools, WK8 current research literature.", False)],
], size=25, space=12)
fill(box["SDG mapping with justification"], "SDG mapping", [
    [("SDG 3 (Target 3.4): ", True), ("well-being for people with ALS, stroke or locked-in syndrome who keep comprehension but lose speech: we measure what non-invasive decoding can deliver.", False)],
    [("SDG 9 (Target 9.5): ", True), ("open, reproducible research capability in neural-interface engineering on free compute.", False)],
    [("SDG 10 (Target 10.2): ", True), ("communication access is a precondition for inclusion of people with severe motor disability.", False)],
    [("Scope: ", True), ("enabling research, not a deployed device; the contribution is evidence of what non-invasive decoding can and cannot achieve.", False)],
], size=25, space=14)

# ---------------------------------------------------------------- references + outcome
fill(box["References"], "References", [
    "[1] S. d'Ascoli et al., TRIBE v2, arXiv:2605.04326, 2026.  [2] J. L. Gardner, T. Liu, eNeuro 6(2), 2019.  [3] J. Tang et al., Nat. Neurosci. 26, 2023.  [4] A. LeBel et al., Sci. Data 10, 555, 2023.  [5] K. N. Kay et al., Nature 452, 2008.  [6] P. C. Hansen, Discrete Inverse Problems, SIAM, 2010.  [7] J. Lévy et al., arXiv:2502.17480, 2025.  [8] C. Destrieux et al., NeuroImage 53, 2010.",
], size=17, head_size=28, space=0)
fill(box["Project outcome"], "Project outcome", [
    [("Complete and reproducible: ", True), ("IEEE-format manuscript (ready for submission), open GitHub repository with code, results, demo video and Colab notebooks, review decks and this poster. Finding: encoding accuracy does not certify decoding; measure recoverability.", False)],
], size=18, head_size=28, space=0)

for k in ["References", "Project outcome"]:
    box[k].text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
prs.save(OUT)
print("saved", OUT)
